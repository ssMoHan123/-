"""
甘蔗病虫害多模态识别系统 - Flask Web应用
"""
import os
import uuid
import json
import torch
import torch.nn as nn
from torchvision import transforms, models
from PIL import Image
from flask import (Flask, render_template, request, jsonify,
                   redirect, url_for, flash, send_file)
from werkzeug.utils import secure_filename
from datetime import datetime

from config import (BASE_DIR, UPLOAD_DIR, BEST_MODEL_PATH, MODEL_DIR,
                    CLASS_NAMES, CLASS_NAMES_CN, TREATMENT_ADVICE,
                    FLASK_CONFIG, TRAIN_CONFIG,
                    DEEPSEEK_CONFIG)
from database import (init_db, insert_recognition_record, get_all_records,
                      get_record_by_id, get_statistics, delete_record)
import cv2
import requests
import tempfile
import shutil
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

VIDEO_CONFIG = {
    'max_video_size': 100 * 1024 * 1024,
    'allowed_extensions': {'mp4', 'avi', 'mov', 'mkv', 'webm'},
    'frame_sample_rate': 10,
}
app = Flask(__name__)
app.config['SECRET_KEY'] = FLASK_CONFIG['SECRET_KEY']
app.config['MAX_CONTENT_LENGTH'] = FLASK_CONFIG['MAX_CONTENT_LENGTH']

# 全局变量：模型和设备
model = None
device = None
class_names = None
transform = None
temperature = 1.0


def load_model():
    """加载训练好的模型"""
    global model, device, class_names, transform, temperature

    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')

    if not os.path.exists(BEST_MODEL_PATH):
        print("[警告] 模型文件不存在，请先运行 train.py 训练模型")
        return False

    # 加载checkpoint
    checkpoint = torch.load(BEST_MODEL_PATH, map_location=device, weights_only=False)
    class_names = checkpoint.get('classes', CLASS_NAMES)
    num_classes = len(class_names)

    # 构建模型
    model = models.mobilenet_v2(weights=None)
    model.classifier = nn.Sequential(
        nn.Dropout(0.3),
        nn.Linear(model.last_channel, 256),
        nn.ReLU(),
        nn.Dropout(0.2),
        nn.Linear(256, num_classes)
    )
    model.load_state_dict(checkpoint['model_state_dict'])
    model = model.to(device)
    model.eval()

    # 预处理变换
    transform = transforms.Compose([
        transforms.Resize((TRAIN_CONFIG['image_size'], TRAIN_CONFIG['image_size'])),
        transforms.ToTensor(),
        transforms.Normalize(mean=[0.485, 0.456, 0.406],
                             std=[0.229, 0.224, 0.225])
    ])

    global temperature
    val_acc = checkpoint.get('val_acc', 0)
    temperature = checkpoint.get('temperature', 1.0)
    print(f"模型加载成功! 验证准确率: {val_acc*100:.2f}%, 温度参数: {temperature:.4f}, 类别: {class_names}")
    return True


def predict_image(image_path):
    """对单张图片进行预测"""
    if model is None:
        return None, 0.0

    image = Image.open(image_path).convert('RGB')
    input_tensor = transform(image).unsqueeze(0).to(device)

    with torch.no_grad():
        outputs = model(input_tensor)
        calibrated = outputs / temperature
        probabilities = torch.softmax(calibrated, dim=1)
        confidence, predicted_idx = torch.max(probabilities, 1)

    predicted_class = class_names[predicted_idx.item()]
    conf = confidence.item()

    # 获取所有类别的概率
    all_probs = {}
    for i, cls_name in enumerate(class_names):
        all_probs[cls_name] = round(probabilities[0][i].item() * 100, 2)

    return predicted_class, conf, all_probs


ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'bmp', 'gif', 'webp'}


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def allowed_video_file(filename):
    """检查视频文件格式"""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in VIDEO_CONFIG['allowed_extensions']


def extract_video_frames(video_path, sample_rate=10):
    """从视频中提取关键帧"""
    cap = cv2.VideoCapture(video_path)
    frames = []
    frame_indices = []
    frame_count = 0

    while True:
        ret, frame = cap.read()
        if not ret:
            break

        if frame_count % sample_rate == 0:
            frames.append(frame)
            frame_indices.append(frame_count)

        frame_count += 1

    cap.release()
    return frames, frame_indices


def predict_video_frame(frame, transform, model, device, class_names):
    """识别视频帧"""
    frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
    pil_image = Image.fromarray(frame_rgb)

    input_tensor = transform(pil_image).unsqueeze(0).to(device)

    with torch.no_grad():
        outputs = model(input_tensor)
        calibrated = outputs / temperature
        probabilities = torch.softmax(calibrated, dim=1)
        confidence, predicted_idx = torch.max(probabilities, 1)

    predicted_class = class_names[predicted_idx.item()]
    confidence_score = confidence.item()

    return predicted_class, confidence_score

# ==================== 路由 ====================

@app.route('/')
def index():
    """首页"""
    stats = get_statistics()
    model_loaded = model is not None
    return render_template('index.html', stats=stats, model_loaded=model_loaded)


@app.route('/recognize', methods=['GET', 'POST'])
def recognize():
    """识别页面"""
    if request.method == 'GET':
        return render_template('recognize.html', model_loaded=(model is not None))

    # POST - 处理上传和识别
    if model is None:
        return jsonify({'error': '模型未加载，请先训练模型'}), 400

    if 'image' not in request.files:
        return jsonify({'error': '未选择图片文件'}), 400

    file = request.files['image']
    if file.filename == '':
        return jsonify({'error': '未选择图片文件'}), 400

    if not allowed_file(file.filename):
        return jsonify({'error': f'不支持的文件格式，请上传 {", ".join(ALLOWED_EXTENSIONS)} 格式的图片'}), 400

    # 保存上传文件
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    ext = file.filename.rsplit('.', 1)[1].lower()
    filename = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:8]}.{ext}"
    filepath = os.path.join(UPLOAD_DIR, filename)
    file.save(filepath)

    # 预测
    try:
        predicted_class, confidence, all_probs = predict_image(filepath)
    except Exception as e:
        return jsonify({'error': f'识别失败: {str(e)}'}), 500

    predicted_class_cn = CLASS_NAMES_CN.get(predicted_class, predicted_class)
    treatment = TREATMENT_ADVICE.get(predicted_class, '暂无防治建议')

    # 获取环境数据
    temperature = request.form.get('temperature', type=float)
    humidity = request.form.get('humidity', type=float)
    rainfall = request.form.get('rainfall', type=float)
    soil_moisture = request.form.get('soil_moisture', type=float)
    notes = request.form.get('notes', '')

    # 保存到数据库
    record_id = insert_recognition_record(
        image_filename=filename,
        image_path=f'uploads/{filename}',
        predicted_class=predicted_class,
        predicted_class_cn=predicted_class_cn,
        confidence=confidence,
        treatment_advice=treatment,
        temperature=temperature,
        humidity=humidity,
        rainfall=rainfall,
        soil_moisture=soil_moisture,
        notes=notes
    )

    result = {
        'record_id': record_id,
        'image_url': f'/static/uploads/{filename}',
        'predicted_class': predicted_class,
        'predicted_class_cn': predicted_class_cn,
        'confidence': round(confidence * 100, 2),
        'treatment_advice': treatment,
        'all_probabilities': all_probs,
        'environment': {
            'temperature': temperature,
            'humidity': humidity,
            'rainfall': rainfall,
            'soil_moisture': soil_moisture,
        }
    }

    return jsonify(result)


@app.route('/recognize_video', methods=['POST'])
def recognize_video():
    """视频文件识别"""
    if model is None:
        return jsonify({'error': '模型未加载，请先训练模型'}), 400

    if 'video' not in request.files:
        return jsonify({'error': '未选择视频文件'}), 400

    file = request.files['video']
    if file.filename == '':
        return jsonify({'error': '未选择视频文件'}), 400

    if not allowed_video_file(file.filename):
        return jsonify({'error': f'不支持的视频格式，请上传 MP4、AVI、MOV、MKV、WEBM 格式'}), 400

    # 保存临时视频文件
    temp_dir = tempfile.mkdtemp()
    video_filename = secure_filename(file.filename)
    video_path = os.path.join(temp_dir, video_filename)
    file.save(video_path)

    try:
        # 提取视频帧
        frames, frame_indices = extract_video_frames(video_path, VIDEO_CONFIG['frame_sample_rate'])

        if len(frames) == 0:
            return jsonify({'error': '无法读取视频帧'}), 400

        # 识别每一帧
        frame_results = []
        class_votes = {}

        for i, frame in enumerate(frames):
            predicted_class, confidence = predict_video_frame(
                frame, transform, model, device, class_names
            )

            frame_results.append({
                'frame_index': frame_indices[i],
                'predicted_class': predicted_class,
                'predicted_class_cn': CLASS_NAMES_CN.get(predicted_class, predicted_class),
                'confidence': round(confidence * 100, 2)
            })

            # 投票统计
            class_votes[predicted_class] = class_votes.get(predicted_class, 0) + 1

        # 综合投票结果
        final_class = max(class_votes, key=class_votes.get)
        total_frames = len(frames)
        # 用投票给最终类别的帧的平均置信度作为综合置信度，而非简单投票率
        winner_confidences = [fr['confidence'] for fr in frame_results
                              if fr['predicted_class'] == final_class]
        confidence_percentage = sum(winner_confidences) / len(winner_confidences)

        # 保存预览图（OpenCV读取为BGR，需转RGB后用PIL保存以保证颜色正确）
        preview_frame = frames[0]
        preview_frame_rgb = cv2.cvtColor(preview_frame, cv2.COLOR_BGR2RGB)
        preview_path = os.path.join(temp_dir, 'preview.jpg')
        Image.fromarray(preview_frame_rgb).save(preview_path)

        # 复制预览图到uploads目录
        preview_filename = f"video_preview_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg"
        final_preview_path = os.path.join(UPLOAD_DIR, preview_filename)
        shutil.copy(preview_path, final_preview_path)

        # 读取环境数据
        temperature = request.form.get('temperature', type=float)
        humidity = request.form.get('humidity', type=float)
        rainfall = request.form.get('rainfall', type=float)
        soil_moisture = request.form.get('soil_moisture', type=float)
        notes = request.form.get('notes', '')
        video_notes = f'视频识别，共分析{total_frames}帧，投票结果：{class_votes}'
        if notes:
            video_notes = notes + ' | ' + video_notes

        # 保存识别记录
        record_id = insert_recognition_record(
            image_filename=video_filename,
            image_path=f'uploads/{preview_filename}',
            predicted_class=final_class,
            predicted_class_cn=CLASS_NAMES_CN.get(final_class, final_class),
            confidence=confidence_percentage / 100,
            treatment_advice=TREATMENT_ADVICE.get(final_class, ''),
            temperature=temperature,
            humidity=humidity,
            rainfall=rainfall,
            soil_moisture=soil_moisture,
            notes=video_notes
        )

        # 清理临时文件
        shutil.rmtree(temp_dir)

        return jsonify({
            'success': True,
            'record_id': record_id,
            'final_class': final_class,
            'final_class_cn': CLASS_NAMES_CN.get(final_class, final_class),
            'confidence': round(confidence_percentage, 2),
            'total_frames': total_frames,
            'frame_results': frame_results[:10],
            'class_votes': class_votes,
            'treatment_advice': TREATMENT_ADVICE.get(final_class, '')
        })

    except Exception as e:
        shutil.rmtree(temp_dir)
        return jsonify({'error': f'视频识别失败: {str(e)}'}), 500

@app.route('/video')
def video_recognize():
    """视频识别页面"""
    return render_template('video.html', model_loaded=(model is not None))



@app.route('/history')
def history():
    """历史记录页面"""
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '')
    class_filter = request.args.get('class_filter', 'all')
    per_page = 12

    records, total = get_all_records(page=page, per_page=per_page,
                                      search=search, class_filter=class_filter)
    total_pages = (total + per_page - 1) // per_page

    return render_template('history.html',
                           records=records,
                           page=page,
                           total_pages=total_pages,
                           total=total,
                           search=search,
                           class_filter=class_filter,
                           class_names=CLASS_NAMES,
                           class_names_cn=CLASS_NAMES_CN)


@app.route('/record/<int:record_id>')
def record_detail(record_id):
    """记录详情"""
    record = get_record_by_id(record_id)
    if not record:
        flash('记录不存在', 'error')
        return redirect(url_for('history'))
    return render_template('record_detail.html', record=record,
                           class_names_cn=CLASS_NAMES_CN)


@app.route('/api/delete_record/<int:record_id>', methods=['POST'])
def api_delete_record(record_id):
    """删除记录API"""
    record = get_record_by_id(record_id)
    if not record:
        return jsonify({'error': '记录不存在'}), 404

    # 删除上传的图片
    image_path = os.path.join(BASE_DIR, 'static', record['image_path'])
    if os.path.exists(image_path):
        os.remove(image_path)

    delete_record(record_id)
    return jsonify({'success': True, 'message': '记录已删除'})


@app.route('/api/export_report/<int:record_id>', methods=['GET', 'POST'])
def export_report(record_id):
    """导出识别报告为Word文档（POST可附带AI分析文本）"""
    record = get_record_by_id(record_id)
    if not record:
        flash('记录不存在', 'error')
        return redirect(url_for('history'))

    ai_analysis = None
    if request.method == 'POST':
        ai_analysis = request.form.get('ai_analysis', '').strip()

    doc = Document()

    # 标题
    title = doc.add_heading('甘蔗病虫害诊断报告', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 基本信息
    doc.add_heading('基本信息', level=2)
    doc.add_paragraph(f'报告编号: #{record["id"]}')
    doc.add_paragraph(f'识别时间: {record["created_at"]}')

    # 诊断结果
    doc.add_heading('诊断结果', level=2)
    disease_badge = '病害' if record['predicted_class'] != 'healthy' else '健康'
    doc.add_paragraph(f'诊断类别: {record["predicted_class_cn"]}（{record["predicted_class"]}）')
    doc.add_paragraph(f'置信度: {round(record["confidence"] * 100, 1)}%')
    doc.add_paragraph(f'诊断结论: {disease_badge}')

    # 插入图片
    image_full_path = os.path.join(BASE_DIR, 'static', record['image_path'])
    if os.path.exists(image_full_path):
        doc.add_heading('诊断图片', level=2)
        try:
            doc.add_picture(image_full_path, width=Inches(4.0))
            last_paragraph = doc.paragraphs[-1]
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            doc.add_paragraph('(图片无法加载)')

    # 防治建议
    if record['treatment_advice']:
        doc.add_heading('防治建议', level=2)
        doc.add_paragraph(record['treatment_advice'])

    # 环境数据
    env_items = []
    if record['temperature'] is not None:
        env_items.append(f'温度: {record["temperature"]}°C')
    if record['humidity'] is not None:
        env_items.append(f'湿度: {record["humidity"]}%')
    if record['rainfall'] is not None:
        env_items.append(f'降雨量: {record["rainfall"]}mm')
    if record['soil_moisture'] is not None:
        env_items.append(f'土壤湿度: {record["soil_moisture"]}%')

    if env_items:
        doc.add_heading('环境数据', level=2)
        for item in env_items:
            doc.add_paragraph(item, style='List Bullet')

    # 备注
    if record['notes']:
        doc.add_heading('备注', level=2)
        doc.add_paragraph(record['notes'])

    # AI 深度分析
    if ai_analysis:
        doc.add_heading('AI 深度分析（DeepSeek）', level=2)
        doc.add_paragraph(ai_analysis)

    # 页脚
    doc.add_paragraph('')
    doc.add_paragraph('—— 本报告由甘蔗病虫害多模态识别系统自动生成 ——').alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 保存到内存
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f'甘蔗病虫害诊断报告_#{record["id"]}_{record["created_at"][:10]}.docx'
    return send_file(
        buffer,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=filename
    )


@app.route('/statistics')
def statistics():
    """统计分析页面"""
    stats = get_statistics()
    return render_template('statistics.html', stats=stats,
                           class_names_cn=CLASS_NAMES_CN)


@app.route('/api/statistics')
def api_statistics():
    """统计数据API"""
    stats = get_statistics()
    return jsonify(stats)


@app.route('/api/deepseek_analysis', methods=['POST'])
def deepseek_analysis():
    """调用DeepSeek大模型进行辅助诊断分析"""
    if not DEEPSEEK_CONFIG['api_key']:
        return jsonify({'error': 'DeepSeek API Key 未配置，请在 config.py 中填写 DEEPSEEK_CONFIG[\'api_key\']'}), 400

    predicted_class = request.form.get('predicted_class', '')
    predicted_class_cn = request.form.get('predicted_class_cn', '')
    confidence = request.form.get('confidence', '')
    all_probs = request.form.get('all_probs', '')
    symptom_text = request.form.get('symptom_text', '').strip()
    temperature = request.form.get('temperature', type=float)
    humidity = request.form.get('humidity', type=float)
    rainfall = request.form.get('rainfall', type=float)
    soil_moisture = request.form.get('soil_moisture', type=float)

    if not symptom_text:
        return jsonify({'error': '请输入症状描述'}), 400

    # 构建环境数据描述
    env_parts = []
    if temperature is not None:
        env_parts.append(f'温度: {temperature}°C')
    if humidity is not None:
        env_parts.append(f'湿度: {humidity}%')
    if rainfall is not None:
        env_parts.append(f'降雨量: {rainfall}mm')
    if soil_moisture is not None:
        env_parts.append(f'土壤湿度: {soil_moisture}%')

    env_desc = '、'.join(env_parts) if env_parts else '无环境数据'
    treatment = TREATMENT_ADVICE.get(predicted_class, '')

    system_prompt = (
        '你是一位经验丰富的甘蔗病虫害诊断专家。'
        '用户会提供以下信息：1) 深度学习视觉模型的识别结果（类别、置信度、各类别概率分布）；'
        '2) 田间环境数据；3) 农户观察到的额外症状描述。'
        '请综合以上多模态信息，给出专业的诊断分析，包括：'
        '诊断结论（视觉模型判断与文字描述是否吻合，不一致时给出你的倾向性判断）、'
        '病害严重程度评估、针对性防治建议、以及需要注意的田间管理要点。'
        '请用中文回答，控制在300字以内，条理清晰。'
    )

    user_prompt = (
        f'【视觉识别结果】\n'
        f'识别类别: {predicted_class_cn}（{predicted_class}）\n'
        f'置信度: {confidence}%\n'
        f'各类别概率: {all_probs}\n\n'
        f'【环境数据】\n{env_desc}\n\n'
        f'【农户补充描述】\n{symptom_text}\n\n'
        f'【系统防治建议参考】\n{treatment}'
    )

    try:
        resp = requests.post(
            DEEPSEEK_CONFIG['api_url'],
            headers={
                'Authorization': f'Bearer {DEEPSEEK_CONFIG["api_key"]}',
                'Content-Type': 'application/json'
            },
            json={
                'model': DEEPSEEK_CONFIG['model'],
                'messages': [
                    {'role': 'system', 'content': system_prompt},
                    {'role': 'user', 'content': user_prompt}
                ],
                'max_tokens': DEEPSEEK_CONFIG['max_tokens'],
                'temperature': DEEPSEEK_CONFIG['temperature']
            },
            timeout=60
        )
        resp.raise_for_status()
        result = resp.json()
        analysis = result['choices'][0]['message']['content']
        return jsonify({'success': True, 'analysis': analysis})

    except requests.exceptions.Timeout:
        return jsonify({'error': 'DeepSeek API 请求超时，请稍后重试'}), 500
    except requests.exceptions.RequestException as e:
        return jsonify({'error': f'DeepSeek API 请求失败: {str(e)}'}), 500


@app.route('/about')
def about():
    """关于页面"""
    model_info = {}
    results_path = os.path.join(MODEL_DIR, 'training_results.json')
    if os.path.exists(results_path):
        with open(results_path, 'r', encoding='utf-8') as f:
            model_info = json.load(f)
    return render_template('about.html', model_info=model_info)


# ==================== 启动 ====================
if __name__ == '__main__':
    print("=" * 60)
    print("甘蔗病虫害多模态识别系统")
    print("=" * 60)

    # 初始化数据库
    init_db()

    # 创建必要目录
    os.makedirs(UPLOAD_DIR, exist_ok=True)

    # 加载模型
    load_model()

    print("\n启动Web服务...")
    print("访问地址: http://127.0.0.1:5000")
    print("=" * 60)
    app.run(host='0.0.0.0', port=5000, debug=True)
